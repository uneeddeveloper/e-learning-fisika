# -*- coding: utf-8 -*-
"""Generator Buku Panduan Aplikasi E-Learning Fisika (.docx)."""

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---------- Palet warna ----------
NAVY = RGBColor(0x1E, 0x3A, 0x8A)      # judul
BLUE = RGBColor(0x25, 0x63, 0xEB)      # aksen
INK = RGBColor(0x1F, 0x29, 0x37)       # teks utama
GREY = RGBColor(0x6B, 0x72, 0x80)      # teks sekunder
LIGHT = RGBColor(0xEF, 0xF4, 0xFF)     # latar tabel header

doc = Document()

# ---------- Style dasar ----------
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.font.color.rgb = INK
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15

# Margin
for s in doc.sections:
    s.top_margin = Inches(0.9)
    s.bottom_margin = Inches(0.9)
    s.left_margin = Inches(1.0)
    s.right_margin = Inches(1.0)


def set_cell_bg(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def add_hrule(paragraph, color="2563EB", size="12"):
    pPr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)
    pPr.append(pbdr)


def heading1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = NAVY
    add_hrule(p)
    return p


def heading2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(13.5)
    r.font.color.rgb = BLUE
    return p


def heading3(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(11.5)
    r.font.color.rgb = INK
    return p


def body(text, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = italic
    return p


def bullet(text, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_lead:
        rb = p.add_run(bold_lead)
        rb.bold = True
        p.add_run(text)
    else:
        p.add_run(text)
    return p


def numbered(text, bold_lead=None):
    p = doc.add_paragraph(style="List Number")
    if bold_lead:
        rb = p.add_run(bold_lead)
        rb.bold = True
        p.add_run(text)
    else:
        p.add_run(text)
    return p


def note_box(label, text, fill="EFF4FF", border="2563EB"):
    """Kotak catatan/tip satu sel."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_bg(cell, fill)
    # border
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "6")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), border)
        borders.append(e)
    tcPr.append(borders)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    rl = p.add_run(label + "  ")
    rl.bold = True
    rl.font.color.rgb = NAVY
    p.add_run(text)
    doc.add_paragraph()
    return tbl


def make_table(headers, rows, widths=None):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = tbl.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_bg(hdr[i], "1E3A8A")
        para = hdr[i].paragraphs[0]
        run = para.add_run(h)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.size = Pt(10.5)
    for row in rows:
        cells = tbl.add_row().cells
        for i, val in enumerate(row):
            para = cells[i].paragraphs[0]
            run = para.add_run(str(val))
            run.font.size = Pt(10.5)
    if widths:
        for ridx, r in enumerate(tbl.rows):
            for i, w in enumerate(widths):
                r.cells[i].width = Inches(w)
    doc.add_paragraph()
    return tbl


def page_break():
    doc.add_page_break()


# ============================================================
# HALAMAN SAMPUL
# ============================================================
for _ in range(4):
    doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("BUKU PANDUAN PENGGUNAAN")
r.bold = True
r.font.size = Pt(16)
r.font.color.rgb = GREY

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("E-LEARNING FISIKA")
r.bold = True
r.font.size = Pt(40)
r.font.color.rgb = NAVY

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_hrule(p, size="18")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Platform Pembelajaran Fisika Interaktif")
r.font.size = Pt(14)
r.font.color.rgb = BLUE
r.italic = True

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Materi  ·  Soal Latihan & Tes  ·  Virtual Lab Simulasi")
r.font.size = Pt(11)
r.font.color.rgb = GREY

for _ in range(8):
    doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Panduan untuk Guru & Siswa")
r.bold = True
r.font.size = Pt(12)
r.font.color.rgb = INK

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Versi 1.0")
r.font.size = Pt(11)
r.font.color.rgb = GREY

page_break()

# ============================================================
# DAFTAR ISI (manual)
# ============================================================
heading1("Daftar Isi")
toc_items = [
    ("1.", "Pendahuluan"),
    ("2.", "Peran Pengguna"),
    ("3.", "Persyaratan & Cara Mengakses"),
    ("4.", "Memulai: Daftar, Login & Logout"),
    ("5.", "Panduan untuk Siswa"),
    ("6.", "Panduan untuk Guru"),
    ("7.", "Katalog Simulasi Virtual Lab"),
    ("8.", "Pertanyaan Umum (FAQ) & Solusi Masalah"),
    ("9.", "Glosarium Istilah"),
]
for num, title in toc_items:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    rn = p.add_run(f"{num}  ")
    rn.bold = True
    rn.font.color.rgb = BLUE
    rt = p.add_run(title)
    rt.font.size = Pt(11.5)

page_break()

# ============================================================
# 1. PENDAHULUAN
# ============================================================
heading1("1. Pendahuluan")

heading2("1.1 Tentang Aplikasi")
body(
    "E-Learning Fisika adalah platform pembelajaran berbasis web yang dirancang untuk "
    "membantu guru menyampaikan materi fisika dan siswa belajar secara mandiri dan interaktif. "
    "Aplikasi menggabungkan tiga komponen belajar dalam satu tempat: materi pembelajaran "
    "(video & bacaan), soal latihan/tes pilihan ganda dengan penilaian otomatis, serta "
    "laboratorium virtual (Virtual Lab) berisi simulasi fisika yang bisa dimainkan langsung "
    "dari browser."
)
body(
    "Aplikasi dapat diakses melalui browser pada komputer maupun perangkat mobile, tanpa perlu "
    "instalasi tambahan."
)

heading2("1.2 Fitur Utama")
bullet("Pengelolaan materi pembelajaran berupa video (YouTube) dan bacaan.", "Materi: ")
bullet(
    "Pembuatan soal pilihan ganda dengan dua mode — Latihan (boleh diulang) dan Tes "
    "(sekali kerjakan) — lengkap dengan penilaian otomatis dan rekap nilai.",
    "Soal Latihan & Tes: ",
)
bullet(
    "Pengaturan batas waktu (deadline) serta status aktif/nonaktif untuk setiap latihan.",
    "Kontrol Waktu: ",
)
bullet(
    "Koleksi simulasi fisika interaktif (PhET dan simulasi custom) yang dilengkapi worksheet.",
    "Virtual Lab: ",
)
bullet(
    "Rekap jawaban worksheet lab dan nilai latihan siswa yang dapat dipantau guru.",
    "Pemantauan: ",
)
bullet(
    "Sistem akun dengan dua peran: Guru dan Siswa.",
    "Autentikasi: ",
)

# ============================================================
# 2. PERAN PENGGUNA
# ============================================================
heading1("2. Peran Pengguna")
body(
    "Aplikasi mengenal dua peran utama. Tampilan menu dan kewenangan menyesuaikan peran akun "
    "yang sedang masuk."
)
make_table(
    ["Peran", "Hak Akses & Kegiatan"],
    [
        ["Guru (Teacher)",
         "Mengakses Teacher Portal: dashboard, kelola materi, membuat & mengatur soal "
         "latihan/tes, melihat nilai siswa, dan meninjau hasil worksheet lab."],
        ["Siswa (Student)",
         "Menjelajahi materi, menonton video & membaca bacaan, mengerjakan latihan/tes, "
         "memainkan simulasi Virtual Lab, dan mengirim worksheet."],
    ],
    widths=[1.6, 4.4],
)
note_box(
    "Catatan:",
    "Katalog materi dan daftar Virtual Lab dapat dijelajahi tanpa login, namun untuk "
    "mengerjakan latihan/tes dan mengirim worksheet, pengguna wajib masuk terlebih dahulu.",
)

# ============================================================
# 3. PERSYARATAN
# ============================================================
heading1("3. Persyaratan & Cara Mengakses")
heading2("3.1 Persyaratan")
bullet("Perangkat: komputer (PC/laptop) atau smartphone/tablet.", "")
bullet("Browser modern terbaru: Google Chrome, Microsoft Edge, Mozilla Firefox, atau Safari.", "")
bullet("Koneksi internet aktif (diperlukan untuk video YouTube dan simulasi PhET).", "")

heading2("3.2 Cara Mengakses")
numbered("Buka browser pada perangkat Anda.", "")
numbered("Masukkan alamat (URL) aplikasi yang diberikan, lalu tekan Enter.", "")
numbered(
    "Halaman awal akan menampilkan katalog Materi Fisika. Dari sini Anda dapat menjelajah "
    "atau masuk ke akun.",
    "",
)

note_box(
    "Akun Demo:",
    "Guru — email: teacher@demo.com, password: teacher123.   |   "
    "Siswa — email: student@demo.com, password: student123.   "
    "(Disarankan mengganti kata sandi setelah aplikasi digunakan secara resmi.)",
    fill="FFF7E6", border="D97706",
)

# ============================================================
# 4. MEMULAI
# ============================================================
heading1("4. Memulai: Daftar, Login & Logout")

heading2("4.1 Mendaftar Akun Baru")
numbered("Pada halaman login, klik tautan “Daftar di sini”.", "")
numbered("Isi Nama lengkap, Email, dan Password (minimal 8 karakter).", "")
numbered("Klik tombol “Daftar”.", "")
numbered("Akun otomatis masuk dan diarahkan ke halaman Materi.", "")
note_box(
    "Penting:",
    "Akun yang dibuat lewat pendaftaran mandiri berperan sebagai Siswa. Akun Guru "
    "disiapkan oleh administrator/pengelola aplikasi.",
)

heading2("4.2 Masuk (Login)")
numbered("Buka halaman Login (tombol “Login” di pojok kanan atas).", "")
numbered("Masukkan Email dan Password yang terdaftar.", "")
numbered("Klik “Sign in”. Setelah berhasil, Anda diarahkan ke halaman Materi.", "")

heading2("4.3 Keluar (Logout)")
body(
    "Klik tombol Logout. Untuk siswa, tombol berada di menu navigasi atas; untuk guru, tombol "
    "berada di kartu profil pada panel kiri (sidebar) Teacher Portal."
)

# ============================================================
# 5. PANDUAN SISWA
# ============================================================
page_break()
heading1("5. Panduan untuk Siswa")
body(
    "Bagian ini menjelaskan cara siswa menggunakan aplikasi: menemukan materi, belajar, "
    "mengerjakan latihan/tes, dan memainkan simulasi lab."
)

heading2("5.1 Menjelajah Materi")
numbered("Buka menu “Materi” pada navigasi atas.", "")
numbered("Daftar materi ditampilkan dalam bentuk kartu. Setiap kartu menampilkan jenis materi (Video, Reading, atau Quiz), judul, dan tanggal unggah.", "")
numbered("Gunakan kotak pencarian untuk mencari materi berdasarkan judul atau topik.", "")
numbered("Klik kartu atau tombol “Buka” untuk membuka materi.", "")

heading2("5.2 Jenis Materi")
heading3("a. Video")
body("Materi video diputar langsung di halaman melalui pemutar YouTube yang tertanam. Klik tombol putar untuk menonton.")
heading3("b. Reading (Bacaan)")
body("Materi bacaan menampilkan teks pembelajaran yang dapat dibaca langsung di halaman.")
heading3("c. Quiz (Latihan/Tes)")
body("Materi bertipe Quiz berisi soal pilihan ganda yang dikerjakan dan dinilai otomatis. Lihat bagian 5.3.")

heading2("5.3 Mengerjakan Latihan / Tes")
numbered("Buka materi bertipe Quiz dari daftar Materi.", "")
numbered(
    "Perhatikan label mode di bagian atas: “Latihan · boleh diulang” atau "
    "“Tes · sekali kerjakan”. Bila ada batas waktu, deadline juga ditampilkan.",
    "",
)
numbered("Pilih satu jawaban untuk setiap soal dengan mengklik pilihan.", "")
numbered("Pastikan semua soal terjawab (indikator “terjawab” menunjukkan progres).", "")
numbered("Klik “Kumpulkan jawaban” untuk mengirim.", "")
numbered("Nilai langsung tampil: skor (0–100) beserta jumlah jawaban benar dari total soal.", "")

heading3("Mode Latihan vs Tes")
make_table(
    ["Mode", "Perilaku"],
    [
        ["Latihan", "Boleh dikerjakan berulang kali. Tombol “Kerjakan ulang” tersedia, dan nilai yang diambil adalah hasil terbaru."],
        ["Tes", "Hanya boleh dikerjakan satu kali. Setelah dikumpulkan, nilai terkunci dan tidak dapat diulang."],
    ],
    widths=[1.4, 4.6],
)
note_box(
    "Tentang Deadline:",
    "Bila guru menetapkan batas waktu, latihan tidak dapat dikerjakan lagi setelah deadline "
    "terlewati. Siswa akan melihat status “Waktu pengerjaan habis”.",
    fill="FDECEC", border="DC2626",
)

heading2("5.4 Virtual Lab & Worksheet")
numbered("Buka menu “Lab” pada navigasi atas.", "")
numbered("Cari simulasi dengan kotak pencarian, atau saring berdasarkan kategori (Mekanika, Gelombang, Listrik & Magnet, Optik, Energi & Termo, Fisika Modern).", "")
numbered("Klik sebuah lab untuk membuka halaman simulasi.", "")
numbered("Mainkan simulasi: ubah variabel (misalnya sudut, kecepatan, panjang tali) dan amati hasilnya. Simulasi PhET juga dapat dibuka layar penuh.", "")
numbered("Baca “Tujuan Pembelajaran” dan “Cara mengerjakan” di panel samping.", "")
numbered("Isi Worksheet: jawab setiap pertanyaan pada kolom yang tersedia. Indikator menunjukkan jumlah soal yang sudah terisi.", "")
numbered("Setelah semua terisi, klik “Kirim worksheet”. Jawaban tersimpan dan terkirim ke guru.", "")
note_box(
    "Tips:",
    "Tombol “Kirim worksheet” baru aktif setelah seluruh pertanyaan terjawab. Anda "
    "dapat menyimpan ulang jawaban kapan saja (“Simpan ulang”), dan tombol "
    "“Reset jawaban” mengosongkan seluruh kolom.",
)

# ============================================================
# 6. PANDUAN GURU
# ============================================================
page_break()
heading1("6. Panduan untuk Guru")
body(
    "Guru bekerja melalui Teacher Portal. Panel navigasi di sisi kiri menyediakan menu: "
    "Dashboard, Student Gallery, Kelola Materi, Soal Latihan, Virtual Lab, dan Hasil Worksheet."
)

heading2("6.1 Dashboard")
body(
    "Dashboard adalah halaman ringkasan yang menampilkan statistik kelas dan daftar materi "
    "terbaru. Gunakan sebagai titik awal sebelum masuk ke menu pengelolaan."
)

heading2("6.2 Kelola Materi")
body("Menu “Kelola Materi” digunakan untuk menambah, mengubah, dan menghapus materi Video & Reading. (Soal Quiz dikelola terpisah di menu Soal Latihan.)")
heading3("Menambah Materi")
numbered("Klik tombol “Tambah Materi”.", "")
numbered("Pilih Tipe: Video atau Reading.", "")
numbered("Isi Judul materi.", "")
numbered("Untuk Reading: tulis isi materi pada kolom Content. Untuk Video: tempel URL YouTube (mendukung format watch, youtu.be, embed, dan shorts) — pratinjau video akan tampil otomatis.", "")
numbered("Klik “Buat materi” untuk menyimpan.", "")
heading3("Mengubah / Menghapus Materi")
bullet("Klik ikon pensil pada kartu materi untuk mengedit, lalu simpan perubahan.", "Edit: ")
bullet("Klik ikon tempat sampah untuk menghapus. Konfirmasi diperlukan sebelum materi dihapus.", "Hapus: ")

heading2("6.3 Soal Latihan & Tes")
heading3("Membuat Latihan Baru")
numbered("Buka menu “Soal Latihan”, lalu klik “Buat Latihan”.", "")
numbered("Isi Judul latihan.", "")
numbered("Pilih Mode: Latihan (boleh diulang, ambil nilai terbaru) atau Tes (sekali kerjakan, nilai terkunci).", "")
numbered("Klik “Buat & isi soal”. Anda akan diarahkan ke halaman pengisian soal.", "")

heading3("Menambahkan Soal")
numbered("Pada tab “Soal”, gunakan form “Tambah Soal” di sisi kanan.", "")
numbered("Tulis Pertanyaan.", "")
numbered("Isi pilihan jawaban (minimal 2, maksimal 6). Klik tombol bundar di samping pilihan untuk menandai jawaban yang benar.", "")
numbered("Klik “Simpan soal”. Soal akan muncul pada daftar di sisi kiri.", "")
numbered("Untuk menghapus soal, klik ikon tempat sampah pada kartu soal (nilai terkait juga ikut terhapus).", "")

heading3("Impor Soal dari Word (.docx)")
body(
    "Selain mengisi satu per satu, guru dapat mengimpor banyak soal sekaligus dari file "
    "Microsoft Word. Tulis soal di Word mengikuti format baku berikut:"
)
# Contoh format (kotak monospace)
tbl = doc.add_table(rows=1, cols=1)
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
cell = tbl.cell(0, 0)
set_cell_bg(cell, "0F172A")
example_lines = [
    "1. Apa satuan SI untuk gaya?",
    "A. Joule",
    "B. Newton",
    "C. Watt",
    "D. Pascal",
    "Jawaban: B",
]
cp = cell.paragraphs[0]
cp.paragraph_format.space_after = Pt(0)
for i, line in enumerate(example_lines):
    if i > 0:
        cp.add_run().add_break()
    run = cp.add_run(line)
    run.font.name = "Consolas"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0xE2, 0xE8, 0xF0)
doc.add_paragraph()

body("Aturan penulisan:")
bullet("Nomor soal diawali “1.” atau “1)”.", "Nomor: ")
bullet("Pilihan jawaban diawali huruf “A.” sampai “E.” (boleh juga memakai tanda kurung, mis. “A)”).", "Pilihan: ")
bullet("Kunci jawaban ditulis pada baris tersendiri: “Jawaban: B”. Kata “Kunci”, “Kunci Jawaban”, atau “Answer” juga diterima.", "Kunci: ")
bullet("Teks pertanyaan boleh ditulis lebih dari satu baris; baris akan disambung otomatis.", "Multi-baris: ")

body("Langkah impor:")
numbered("Buka latihan, lalu pada panel “Tambah Soal” klik tombol “Impor dari Word (.docx)”.", "")
numbered("Pilih file .docx dari komputer Anda.", "")
numbered("Periksa pratinjau: soal yang valid bertanda hijau, soal bermasalah bertanda merah beserta keterangan masalahnya.", "")
numbered("Klik “Simpan … soal”. Hanya soal valid yang disimpan; soal bermasalah dilewati.", "")

note_box(
    "Keterbatasan:",
    "File harus berformat .docx (bukan .doc lama). Rumus atau simbol fisika yang dibuat "
    "sebagai gambar / Equation tidak ikut terbaca — tulis sebagai teks biasa (mis. m/s^2) "
    "agar dapat diimpor.",
    fill="FFF7E6", border="D97706",
)

heading3("Pengaturan: Aktif/Nonaktif & Deadline")
bullet("Klik tombol status untuk mengaktifkan atau menonaktifkan latihan. Latihan nonaktif tidak dapat dikerjakan siswa.", "Aktif/Nonaktif: ")
bullet("Tentukan tanggal & jam pada kolom deadline, lalu klik “Simpan deadline”. Gunakan “Hapus deadline” untuk meniadakan batas waktu.", "Deadline: ")
note_box(
    "Cara kerja deadline:",
    "Setelah deadline terlewati, latihan terkunci selama 1 hari (siswa melihat “waktu "
    "pengerjaan habis”), kemudian otomatis berstatus nonaktif.",
    fill="FFF7E6", border="D97706",
)

heading3("Melihat Nilai Siswa")
numbered("Pada halaman latihan, buka tab “Nilai”.", "")
numbered("Tabel menampilkan nama & email siswa, jumlah jawaban benar, skor, dan waktu pengerjaan.", "")
numbered("Klik “Refresh” untuk memuat data terbaru.", "")

heading2("6.4 Virtual Lab")
body(
    "Guru dapat membuka menu “Virtual Lab” untuk melihat seluruh simulasi yang tersedia "
    "bagi siswa, sama seperti tampilan siswa. Gunakan untuk menyiapkan instruksi atau menentukan "
    "lab mana yang akan ditugaskan."
)

heading2("6.5 Hasil Worksheet")
numbered("Buka menu “Hasil Worksheet”.", "")
numbered("Daftar lab yang sudah menerima worksheet ditampilkan beserta jumlah siswa pengirim dan waktu terakhir.", "")
numbered("Klik salah satu lab untuk melihat detail jawaban worksheet per siswa.", "")
numbered("Klik “Refresh” untuk memuat kiriman terbaru.", "")

# ============================================================
# 7. KATALOG LAB
# ============================================================
page_break()
heading1("7. Katalog Simulasi Virtual Lab")
body(
    "Berikut daftar simulasi yang tersedia. Setiap lab dilengkapi tujuan pembelajaran dan "
    "worksheet. Sebagian besar simulasi disediakan oleh PhET Interactive Simulations "
    "(University of Colorado Boulder); satu simulasi bersifat custom (dibuat khusus)."
)
make_table(
    ["Simulasi", "Kategori", "Jenjang", "Durasi"],
    [
        ["Gerak Parabola Interaktif (custom)", "Mekanika", "Kelas X", "20 menit"],
        ["Gaya & Gerak (Forces and Motion)", "Mekanika", "Kelas X", "25 menit"],
        ["Bandul Sederhana (Pendulum Lab)", "Mekanika", "Kelas XI", "20 menit"],
        ["Gelombang pada Tali", "Gelombang", "Kelas XI", "20 menit"],
        ["Interferensi Gelombang", "Gelombang", "Kelas XI", "25 menit"],
        ["Membangun Rangkaian Listrik", "Listrik & Magnet", "Kelas XII", "30 menit"],
        ["Magnet & Kompas", "Listrik & Magnet", "Kelas XII", "15 menit"],
        ["Pembiasan Cahaya", "Optik", "Kelas XI", "20 menit"],
        ["Skate Park: Energi Mekanik", "Energi & Termo", "Kelas X", "20 menit"],
        ["Bangun Atom", "Fisika Modern", "Kelas XII", "15 menit"],
    ],
    widths=[3.0, 1.6, 1.0, 0.9],
)

# ============================================================
# 8. FAQ
# ============================================================
heading1("8. Pertanyaan Umum (FAQ) & Solusi Masalah")

faqs = [
    ("Saya tidak bisa mengerjakan latihan, statusnya “Waktu pengerjaan habis”.",
     "Deadline latihan sudah terlewati. Hubungi guru bila Anda merasa perlu kesempatan tambahan; "
     "guru dapat mengaktifkan kembali atau mengubah deadline."),
    ("Tombol “Kirim worksheet” tidak bisa diklik.",
     "Pastikan seluruh pertanyaan worksheet sudah terisi. Pastikan juga Anda sudah login sebagai siswa."),
    ("Video tidak tampil / muncul peringatan format tidak didukung.",
     "Pastikan URL adalah tautan YouTube yang valid (format watch, youtu.be, embed, atau shorts) "
     "dan koneksi internet aktif."),
    ("Simulasi PhET tidak muncul.",
     "Simulasi PhET dimuat dari internet. Periksa koneksi, atau gunakan tombol “Buka full-screen” "
     "untuk membuka simulasi di tab baru."),
    ("Saya guru tetapi tidak melihat menu Teacher.",
     "Menu Teacher hanya muncul untuk akun berperan Guru. Pastikan Anda masuk dengan akun guru."),
    ("Apakah nilai Tes bisa dikerjakan ulang?",
     "Tidak. Mode Tes hanya dapat dikerjakan satu kali dan nilainya terkunci. Gunakan mode Latihan "
     "bila ingin siswa dapat mengulang."),
    ("Saat impor dari Word, sebagian soal tidak tersimpan / ditandai merah.",
     "Soal tersebut belum sesuai format. Pastikan ada baris “Jawaban: …”, minimal 2 pilihan, dan "
     "huruf kunci jawaban benar-benar ada di antara pilihan. Perbaiki di Word lalu unggah ulang."),
    ("File Word saya tidak bisa diimpor.",
     "Pastikan file berformat .docx (bukan .doc lama). Jika masih gagal, buka di Word lalu "
     "“Save As” ke format .docx, dan hindari menaruh soal di dalam tabel atau sebagai gambar."),
]
for q, a in faqs:
    heading3("T: " + q)
    p = doc.add_paragraph()
    rj = p.add_run("J: ")
    rj.bold = True
    rj.font.color.rgb = BLUE
    p.add_run(a)

# ============================================================
# 9. GLOSARIUM
# ============================================================
heading1("9. Glosarium Istilah")
make_table(
    ["Istilah", "Penjelasan"],
    [
        ["Materi (Lesson)", "Satuan konten pembelajaran: Video, Reading (bacaan), atau Quiz."],
        ["Latihan", "Soal pilihan ganda mode boleh diulang; nilai yang diambil adalah hasil terbaru."],
        ["Tes", "Soal pilihan ganda mode sekali kerjakan; nilai terkunci setelah dikumpulkan."],
        ["Deadline", "Batas waktu pengerjaan latihan/tes yang ditetapkan guru."],
        ["Virtual Lab", "Laboratorium simulasi fisika interaktif yang dijalankan di browser."],
        ["Worksheet", "Lembar pertanyaan yang diisi siswa setelah memainkan simulasi lab."],
        ["PhET", "Penyedia simulasi sains gratis dari University of Colorado Boulder."],
        ["Teacher Portal", "Area kerja khusus guru untuk mengelola materi, soal, dan hasil belajar siswa."],
    ],
    widths=[1.8, 4.2],
)

doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_hrule(p)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("— Akhir Buku Panduan —")
r.italic = True
r.font.color.rgb = GREY

# ---------- Footer ----------
section = doc.sections[0]
footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = fp.add_run("E-Learning Fisika — Buku Panduan Penggunaan")
fr.font.size = Pt(8)
fr.font.color.rgb = GREY

out = r"d:\e-learning-fisika\Buku Panduan E-Learning Fisika.docx"
doc.save(out)
print("Saved:", out)
